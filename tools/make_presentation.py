"""
tools/make_presentation.py - Generator for the 3-slide hackathon presentation.
Outputs submission/711524BCS049_GURU PRASATH C.pptx
"""

import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

METRICS_PATH = 'results/metrics.json'
PPTX_PATH = os.path.join('submission', '711524BCS049_GURU PRASATH C.pptx')

def create_presentation():
    os.makedirs('submission', exist_ok=True)

    # Load evaluated metrics if available
    metrics = {}
    if os.path.exists(METRICS_PATH):
        try:
            with open(METRICS_PATH, 'r') as f:
                metrics = json.load(f)
        except Exception:
            pass

    aug_metrics = metrics.get('night_augmented_model', {})
    clean_acc = aug_metrics.get('clean_test', {}).get('overall_accuracy', None)
    night_acc = aug_metrics.get('simulated_night_test', {}).get('overall_accuracy', None)
    ele_recall = aug_metrics.get('clean_test', {}).get('elephant_recall', None)

    clean_acc_str = f"{clean_acc * 100:.1f}%" if clean_acc is not None else "[Evaluated Accuracy]"
    night_acc_str = f"{night_acc * 100:.1f}%" if night_acc is not None else "[Evaluated Night Acc]"
    ele_recall_str = f"{ele_recall * 100:.1f}%" if ele_recall is not None else "[Evaluated Recall]"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    DARK_BG = RGBColor(15, 23, 42)
    TEXT_LIGHT = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    ACCENT_BLUE = RGBColor(56, 189, 248)
    ACCENT_RED = RGBColor(239, 68, 68)
    CARD_BG = RGBColor(30, 41, 59)

    def add_slide_header(slide, title_text, category_text):
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

        # Category / Subtitle
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

    # --- SLIDE 1: Problem Statement ---
    slide1 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide1, "GajaAlert AI: Human-Elephant Conflict Early Warning System", "Slide 1: Problem Statement & Context")

    # Card 1: Problem & Fatality Context
    c1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = ACCENT_RED

    tf1 = c1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "CRITICAL PROBLEM CONTEXT"
    p1.font.bold = True
    p1.font.size = Pt(16)
    p1.font.color.rgb = ACCENT_RED

    bullets1 = [
        "Human-Elephant Conflicts (HEC) present severe threats to lives and agricultural livelihoods at forest-village borders near Coimbatore, Tamil Nadu.",
        "TN Government Source (2015 - July 2025): Recorded 685 human deaths from wild animal attacks; 522 deaths (76.2%) were caused specifically by elephants.",
        "Existing solutions depend on manual human informants (Valparai SMS) or costly monitored camera towers, introducing delayed warnings."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    # Card 2: Technical Gap & Solution
    c2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = ACCENT_BLUE

    tf2 = c2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "OUR GAJAALERT AI SOLUTION"
    p2.font.bold = True
    p2.font.size = Pt(16)
    p2.font.color.rgb = ACCENT_BLUE

    bullets2 = [
        "Automated Edge Warning: Camera frames are classified automatically via deep learning to broadcast instant alerts.",
        "Low-Light Robustness: Dynamic night-augmentation (brightness U(0.15,0.60), noise, IR tint) optimizes low-light detection.",
        "Explainable & Bilingual: Grad-CAM neural heatmaps explain predictions; instant English & Tamil alerts inform local communities.",
        "Honesty & Scope: Camera feeds & SMS delivery are simulated for prototype verification."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT


    # --- SLIDE 2: Methods Related to the Subjects ---
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide2, "Integrated Technical Stack & Methodologies", "Slide 2: Methods Related to the Subjects (MWP & DNN)")

    # Column 1: MWP Methods
    col1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    col1.fill.solid()
    col1.fill.fore_color.rgb = CARD_BG
    col1.line.color.rgb = ACCENT_BLUE

    tf_mwp = col1.text_frame
    tf_mwp.word_wrap = True
    pm = tf_mwp.paragraphs[0]
    pm.text = "MODERN WEB PROGRAMMING (MWP)"
    pm.font.bold = True
    pm.font.size = Pt(16)
    pm.font.color.rgb = ACCENT_BLUE

    mwp_items = [
        "Frontend: HTML5 Semantic Structure, Custom CSS3 Dark Palette, Vanilla JS ES6+ with Fetch API.",
        "Offline Vendor Architecture: Local vendored Leaflet Map, Chart.js, and font assets under static/vendor/.",
        "REST API Backend: Python Flask REST endpoints (/predict, /api/events, /api/stats, /health).",
        "Persistence & Security: SQLite database, Pillow image validation, 8 MB payload cap, UUID filenames.",
        "Web Audio & TTS: Browser-native Web Audio siren synthesizer & Web Speech API bilingual TTS."
    ]
    for item in mwp_items:
        p = tf_mwp.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    # Column 2: DNN Methods
    col2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    col2.fill.solid()
    col2.fill.fore_color.rgb = CARD_BG
    col2.line.color.rgb = RGBColor(16, 185, 129)

    tf_dnn = col2.text_frame
    tf_dnn.word_wrap = True
    pd = tf_dnn.paragraphs[0]
    pd.text = "DEEP NEURAL NETWORKS (DNN)"
    pd.font.bold = True
    pd.font.size = Pt(16)
    pd.font.color.rgb = RGBColor(16, 185, 129)

    dnn_items = [
        "Architecture: MobileNetV2 Transfer Learning pre-trained on ImageNet with Dense Head (128 units, Dropout 0.3, Softmax).",
        "Two-Phase Training: Phase A frozen backbone (Adam 1e-3); Phase B fine-tuning last 30 layers (Adam 1e-5, BatchNorm frozen).",
        "Custom tf.data Night Augmentation: 40% training frames undergo brightness scaling, IR tinting, and Gaussian noise.",
        "Grad-CAM Explainability: Gradient tape targeting final conv layer (out_relu) overlaid on input frames.",
        "Deterministic Risk Logic: NON-ELEPHANT -> NONE; ELEPHANT -> Base MEDIUM/LOW; +1 Night; +1 Repeat (30m)."
    ]
    for item in dnn_items:
        p = tf_dnn.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT


    # --- SLIDE 3: Working Model Explanation ---
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide3, "End-to-End System Workflow & Empirical Performance", "Slide 3: Working Model Explanation")

    # Workflow Box (Top)
    wf_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.2))
    wf_box.fill.solid()
    wf_box.fill.fore_color.rgb = CARD_BG
    wf_box.line.color.rgb = ACCENT_BLUE

    tf_wf = wf_box.text_frame
    tf_wf.word_wrap = True
    pw = tf_wf.paragraphs[0]
    pw.text = "SYSTEM WORKFLOW: INPUT -> PROCESSING -> DNN -> OUTPUT"
    pw.font.bold = True
    pw.font.size = Pt(15)
    pw.font.color.rgb = ACCENT_BLUE

    steps = [
        "1. INPUT: Drag & drop camera trap image + virtual sensor location select + time mode simulation.",
        "2. PROCESSING: Pillow validates & resizes to 224x224x3; REST payload parsed by Flask service.",
        "3. DNN INFERENCE: MobileNetV2 evaluates class probabilities; Grad-CAM generates out_relu heatmap.",
        "4. RISK & OUTPUT: Safety engine evaluates time/repeat context -> updates Verdict Band -> triggers siren & bilingual alerts."
    ]
    for step in steps:
        p = tf_wf.add_paragraph()
        p.text = step
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    # Metrics Box (Bottom Left)
    mb = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(5.6), Inches(2.5))
    mb.fill.solid()
    mb.fill.fore_color.rgb = CARD_BG
    mb.line.color.rgb = RGBColor(16, 185, 129)

    tf_m = mb.text_frame
    tf_m.word_wrap = True
    pm = tf_m.paragraphs[0]
    pm.text = "MEASURED MODEL EVALUATION"
    pm.font.bold = True
    pm.font.size = Pt(14)
    pm.font.color.rgb = RGBColor(16, 185, 129)

    m_lines = [
        f"• Clean Held-Out Test Accuracy: {clean_acc_str}",
        f"• Simulated Night Test Accuracy: {night_acc_str}",
        f"• Elephant Detection Recall: {ele_recall_str}",
        "• Verification: Evaluated on held-out test split via evaluate.py."
    ]
    for ml in m_lines:
        p = tf_m.add_paragraph()
        p.text = ml
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    # Operational Impact Box (Bottom Right)
    ob = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.2), Inches(5.7), Inches(2.5))
    ob.fill.solid()
    ob.fill.fore_color.rgb = CARD_BG
    ob.line.color.rgb = RGBColor(245, 158, 11)

    tf_o = ob.text_frame
    tf_o.word_wrap = True
    po = tf_o.paragraphs[0]
    po.text = "OPERATIONAL & DEMO HIGHLIGHTS"
    po.font.bold = True
    po.font.size = Pt(14)
    po.font.color.rgb = RGBColor(245, 158, 11)

    o_lines = [
        "• High Recall Priority: Minimizes missed elephant intrusions near village borders.",
        "• Dual Accessibility: Dynamic color-coded Verdict Band + aria-live screen reader support.",
        "• Interactive Dashboard: Real-time Leaflet sensor map, Chart.js sightings graph, and event logs."
    ]
    for ol in o_lines:
        p = tf_o.add_paragraph()
        p.text = ol
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    prs.save(PPTX_PATH)
    print(f"[make_presentation.py] Successfully generated presentation at {PPTX_PATH}")

if __name__ == '__main__':
    create_presentation()
